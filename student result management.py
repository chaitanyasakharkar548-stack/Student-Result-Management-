import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
import os

file = "student_results.xlsx"

# Create Excel file
if not os.path.exists(file):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Roll No", "Class", "Sub1", "Sub2", "Sub3",
               "Sub4", "Sub5", "Total", "Percentage", "Result"])
    wb.save(file)


# Save Student
def save_student():
    name = name_entry.get()
    roll = roll_entry.get()
    cls = class_entry.get()

    try:
        marks = [
            int(sub1.get()),
            int(sub2.get()),
            int(sub3.get()),
            int(sub4.get()),
            int(sub5.get())
        ]
    except:
        messagebox.showerror("Error", "Enter valid marks")
        return

    total = sum(marks)
    percentage = total / 5

    if all(mark >= 35 for mark in marks):
        result = "Pass"
    else:
        result = "Fail"

    wb = load_workbook(file)
    ws = wb.active

    ws.append([
        name, roll, cls,
        marks[0], marks[1], marks[2],
        marks[3], marks[4],
        total, percentage, result
    ])

    wb.save(file)

    messagebox.showinfo("Success", "Student saved successfully")


# Get Result
def get_result():
    roll = search_entry.get()

    wb = load_workbook(file)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):

        if str(row[1]) == roll:

            messagebox.showinfo(
                "Student Result",
                "Name: " + str(row[0]) +
                "\nRoll No: " + str(row[1]) +
                "\nClass: " + str(row[2]) +
                "\nTotal: " + str(row[8]) +
                "\nPercentage: " + str(row[9]) + "%" +
                "\nResult: " + str(row[10])
            )

            wb.close()
            return

    wb.close()

    messagebox.showerror(
        "Error",
        "Student record not found"
    )


# Main Window
root = tk.Tk()
root.title("Student Result Management")
root.geometry("500x600")


tk.Label(
    root,
    text="STUDENT RESULT MANAGEMENT",
    font=("Arial", 18, "bold")
).pack(pady=20)


# Student details
tk.Label(root, text="Name").pack()
name_entry = tk.Entry(root)
name_entry.pack()

tk.Label(root, text="Roll No").pack()
roll_entry = tk.Entry(root)
roll_entry.pack()

tk.Label(root, text="Class").pack()
class_entry = tk.Entry(root)
class_entry.pack()


# Subjects
tk.Label(root, text="Subject 1").pack()
sub1 = tk.Entry(root)
sub1.pack()

tk.Label(root, text="Subject 2").pack()
sub2 = tk.Entry(root)
sub2.pack()

tk.Label(root, text="Subject 3").pack()
sub3 = tk.Entry(root)
sub3.pack()

tk.Label(root, text="Subject 4").pack()
sub4 = tk.Entry(root)
sub4.pack()

tk.Label(root, text="Subject 5").pack()
sub5 = tk.Entry(root)
sub5.pack()


# Save button
tk.Button(
    root,
    text="Save Student",
    command=save_student
).pack(pady=15)


# Search
tk.Label(
    root,
    text="Enter Roll No. to Get Result"
).pack()

search_entry = tk.Entry(root)
search_entry.pack()

tk.Button(
    root,
    text="Get Result",
    command=get_result
).pack(pady=10)


# Exit
tk.Button(
    root,
    text="Exit",
    command=root.destroy
).pack(pady=10)


root.mainloop()
